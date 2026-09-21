"""Bounded document text/table extraction from explicitly allowed local roots."""

from __future__ import annotations

import asyncio
import csv
import io
import zipfile
from collections.abc import Generator, Sequence
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from typing import Annotated, Any

from pydantic import Field

from protolink.tools.builtins._integration import integration_tool
from protolink.tools.builtins.filesystem import FilesystemResource
from protolink.tools.prepared import PreparedTool

_TEXT = {".txt", ".md", ".rst", ".log", ".json", ".yaml", ".yml", ".xml"}


def _archive(data: bytes, limit: int) -> None:
    """Check declared decompressed package size before optional Office parsers run."""
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        members = package.infolist()
        if len(members) > 10000 or sum(member.file_size for member in members) > limit:
            raise ValueError("Document archive exceeds the configured expansion limit")
        if any(member.flag_bits & 1 for member in members):
            raise ValueError("Encrypted document archives are not supported")


def _sections(data: bytes, suffix: str, archive_limit: int) -> Generator[dict[str, Any], None, None]:
    """Yield text sections or table rows with source locations; never execute document content."""
    if suffix in _TEXT:
        if b"\x00" in data:
            raise ValueError("Text documents must contain UTF-8 text without NUL")
        yield {"kind": "text", "location": {"section": 1}, "text": data.decode("utf-8-sig")}
    elif suffix in {".csv", ".tsv"}:
        for row_number, cells in enumerate(
            csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter="\t" if suffix == ".tsv" else ","), 1
        ):
            yield {
                "kind": "table_row",
                "location": {"row": row_number},
                "cells": cells[:100],
                "truncated": len(cells) > 100,
            }
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ImportError(
                "PDF extraction requires pip install 'protolink[documents]' or 'protolink[rag-pdf]'"
            ) from exc
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported")
        for number, page in enumerate(reader.pages, 1):
            yield {"kind": "text", "location": {"page": number}, "text": page.extract_text() or ""}
    elif suffix == ".docx":
        _archive(data, archive_limit)
        try:
            from docx import Document
            from docx.table import Table
        except ImportError as exc:
            raise ImportError("Word extraction requires pip install 'protolink[documents]'") from exc
        document = Document(io.BytesIO(data))
        for index, block in enumerate(document.iter_inner_content(), 1):
            if isinstance(block, Table):
                for number, row in enumerate(block.rows, 1):
                    cells = row.cells
                    yield {
                        "kind": "table_row",
                        "location": {"block": index, "row": number},
                        "cells": [cell.text for cell in cells[:100]],
                        "truncated": len(cells) > 100,
                    }
            else:
                yield {"kind": "text", "location": {"block": index}, "text": block.text}
    elif suffix == ".xlsx":
        _archive(data, archive_limit)
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("Spreadsheet extraction requires pip install 'protolink[documents]'") from exc
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                for number, row in enumerate(
                    sheet.iter_rows(max_col=min(sheet.max_column or 100, 100), values_only=True), 1
                ):
                    cells = [
                        value.isoformat()
                        if isinstance(value, (date, datetime))
                        else ""
                        if value is None
                        else str(value)
                        for value in row
                    ]
                    yield {
                        "kind": "table_row",
                        "location": {"sheet": sheet.title, "row": number},
                        "cells": cells,
                        "truncated": (sheet.max_column or 0) > 100,
                    }
        finally:
            workbook.close()
    else:
        raise ValueError("Unsupported document format; use UTF-8 text, CSV/TSV, PDF, DOCX, or XLSX")


def document_tools(
    *,
    roots: Sequence[str | Path],
    max_file_bytes: int = 8 * 1024 * 1024,
    max_output_chars: int = 20000,
    max_sections: int = 100,
    max_archive_bytes: int = 32 * 1024 * 1024,
    timeout_seconds: float = 30,
) -> tuple[PreparedTool, ...]:
    """Expose read_document/search_document with ``filesystem.read`` capability.

    Args:
        roots: Explicit POSIX directory roots; absolute paths only, no symlinks.
        max_file_bytes: Maximum input file bytes, default 8 MiB.
        max_output_chars: Total extracted text/cell characters, default 20,000.
        max_sections: Maximum sections or table rows, default 100.
        max_archive_bytes: Declared expanded DOCX/XLSX package bytes, default 32 MiB.
        timeout_seconds: Overall await limit, further bounded by run budgets.

    UTF-8 text/CSV/TSV use the standard library. PDF/DOCX/XLSX require the optional
    ``documents`` extra. Results contain file metadata and located text sections or
    table rows; cells are strings and tables are limited to 100 columns. Search is
    literal and covers only the bounded extraction, with truncation reported.
    PDF extraction has no OCR or table reconstruction. Word covers top-level body
    paragraphs/tables; spreadsheets read cached values without evaluating formulas.
    Macros and remote links are never executed. Parsers run on a worker; cancellation
    stops waiting but cannot forcibly stop an in-progress parser. Input/output caps
    are not a hard parser memory limit; use process isolation for hostile documents.
    """
    for limit in (max_output_chars, max_sections, max_archive_bytes):
        if isinstance(limit, bool) or not isinstance(limit, int) or limit < 1:
            raise ValueError("Document limits must be positive integers")
    resource = FilesystemResource(roots, max_file_bytes=max_file_bytes)

    def extract(path: str) -> dict[str, Any]:
        snapshot = resource.read(path)
        data = snapshot.data
        if data is None:
            raise FileNotFoundError(path)
        suffix = Path(path).suffix.lower()
        remaining = max_output_chars
        sections = []
        truncated = False
        with closing(_sections(data, suffix, max_archive_bytes)) as source:
            for section in source:
                if len(sections) == max_sections or remaining == 0:
                    truncated = True
                    break
                if section.get("truncated"):
                    truncated = True
                values = section["cells"] if "cells" in section else [section["text"]]
                clipped = []
                for value in values:
                    clipped.append(value[:remaining])
                    if len(value) > remaining:
                        truncated = True
                        section["truncated"] = True
                    remaining -= len(clipped[-1])
                section["cells" if "cells" in section else "text"] = clipped if "cells" in section else clipped[0]
                sections.append(section)
        return {
            "metadata": {"path": snapshot.revision.resource_id, "format": suffix.lstrip("."), "size_bytes": len(data)},
            "sections": sections,
            "truncated": truncated,
        }

    async def read_document(path: str) -> dict[str, Any]:
        """Extract text, table rows, and source locations. truncated=true means incomplete extraction."""
        return await asyncio.to_thread(extract, path)

    async def search_document(
        path: str, query: str, max_results: Annotated[int, Field(ge=1, le=100)] = 20
    ) -> dict[str, Any]:
        """Search literal text case-insensitively within bounded extraction, returning located snippets."""
        document = await asyncio.to_thread(extract, path)
        matches = []
        limited = False
        for section in document["sections"]:
            text = section.get("text", "\t".join(section.get("cells", [])))
            # Search original offsets so Unicode case folding cannot shift snippets.
            for line_number, line in enumerate(text.splitlines(), 1):
                if query.casefold() not in line.casefold():
                    continue
                if len(matches) == max_results:
                    limited = True
                    break
                matches.append(
                    {
                        "location": {**section["location"], "line": line_number},
                        "text": line[:500],
                        "text_truncated": len(line) > 500,
                    }
                )
            if limited:
                break
        return {"metadata": document["metadata"], "matches": matches, "truncated": document["truncated"] or limited}

    def validate(arguments: dict[str, Any]) -> None:
        resource._target(arguments["path"])
        if "query" in arguments and not arguments["query"]:
            raise ValueError("query must not be empty")

    return tuple(
        integration_tool(
            function,
            capability="filesystem.read",
            target="configured document roots",
            timeout_seconds=timeout_seconds,
            validate=validate,
        )
        for function in (read_document, search_document)
    )
